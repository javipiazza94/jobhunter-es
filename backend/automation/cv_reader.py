"""
cv_reader.py — Extrae el texto de un CV (PDF o DOCX) en local, sin enviarlo a ningún servicio.

PDF: reutiliza `verify_pdf.extract_text_layer` (pypdf y, si no basta, pdftotext).
DOCX: párrafos y tablas con python-docx.
Cualquier problema (tipo no soportado, cifrado, vacío, ilegible) se reporta como `CVReadError`
con un mensaje apto para mostrar al usuario.
"""
import zipfile
from pathlib import Path

from automation.verify_pdf import VerificationError, extract_text_layer


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_PDF_PAGES = 30
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_TEXT_CHARS = 40_000


class CVReadError(Exception):
    """El CV no se pudo leer; el mensaje es legible por el usuario."""


def _pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        # Cifrado solo con contraseña de propietario: se abre con clave de usuario vacía
        if reader.is_encrypted and not reader.decrypt(""):
            raise CVReadError("El PDF está cifrado; expórtalo sin contraseña e inténtalo de nuevo")
    except CVReadError:
        raise
    except Exception:
        pass  # ilegible para pypdf: lo intenta el extractor de reserva y, si falla, se informa abajo
    try:
        text, _pages, _extractor = extract_text_layer(path, max_pages=MAX_PDF_PAGES)
    except (VerificationError, OSError) as exc:
        raise CVReadError("No se pudo leer el PDF (¿está dañado?)") from exc
    return text


def _docx_text(path: Path) -> str:
    try:
        with zipfile.ZipFile(path) as zf:
            if sum(i.file_size for i in zf.infolist()) > MAX_DOCX_UNCOMPRESSED:
                raise CVReadError("El DOCX es demasiado grande al descomprimirse")
    except CVReadError:
        raise
    except Exception as exc:
        raise CVReadError("No se pudo leer el DOCX (¿está dañado?)") from exc
    try:
        from docx import Document

        doc = Document(str(path))
    except Exception as exc:
        raise CVReadError("No se pudo leer el DOCX (¿está dañado?)") from exc
    parts = [p.text for p in doc.paragraphs]
    for section in doc.sections:  # cabeceras y pies (a menudo llevan el nombre y el contacto)
        for part in (section.header, section.footer):
            for p in part.paragraphs:
                if p.text.strip() and p.text not in parts:
                    parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            seen: set[int] = set()
            cells = []
            for cell in row.cells:
                if id(cell._tc) in seen:  # celdas combinadas se repiten
                    continue
                seen.add(id(cell._tc))
                cells.append(cell.text)
            parts.append(" | ".join(c for c in cells if c.strip()))
    return "\n".join(parts)


def extract_text(path: Path) -> str:
    """Texto del CV en `path` (.pdf o .docx); `CVReadError` si no se puede leer o está vacío."""
    path = Path(path)
    if not path.is_file():
        raise CVReadError("No se encontró el archivo del CV")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise CVReadError("El archivo supera el tamaño máximo de 10 MB")
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _pdf_text(path)
    elif suffix == ".docx":
        text = _docx_text(path)
    else:
        raise CVReadError(f"Tipo de archivo no soportado ({suffix or 'sin extensión'}): usa PDF o DOCX")
    if not text.strip():
        raise CVReadError(
            "El archivo está vacío o no contiene texto (¿es un PDF escaneado?); rellena el perfil a mano"
        )
    return text[:MAX_TEXT_CHARS]
